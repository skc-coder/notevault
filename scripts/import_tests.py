import re
import os
import openpyxl

EXCEL_PATH = '/home/skc/Documents/notevault/Test lists.xlsx'
VAULT_PATH = '/home/skc/Documents/notevault'

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

folder_map = {
    'GATE CSE PYQs': 'Tests/GATE CSE PYQs',
    'CSE quiz': 'Tests/CSE Quizzes',
    'DA quiz': 'Tests/DA Quizzes',
    'GO Test Series (129)': 'Tests/GO Test Series',
    'GO Test Series DA (88)': 'Tests/GO Test Series DA',
    'GATE Overflow 2027 (29)': 'Tests/GATE Overflow 2027'
}

def clean_filename(name):
    name = re.sub(r'[\\/*?:"<>|]', '_', name)
    name = re.sub(r'\s+', ' ', name).strip()
    return name

subject_moc_map = {
    'digital logic': 'notes/moc dl.md',
    'dl': 'notes/moc dl.md',
    'co and architecture': 'notes/moc coa.md',
    'computer organization & architecture': 'notes/moc coa.md',
    'operating system': 'notes/moc os.md',
    'operating systems': 'notes/moc os.md',
    'theory of computation': 'notes/moc toc.md',
    'toc': 'notes/moc toc.md',
    'compiler design': 'notes/moc cd.md',
    'cd': 'notes/moc cd.md',
    'databases': 'notes/moc dbms.md',
    'dbms': 'notes/moc dbms.md',
    'computer networks': 'notes/moc cn.md',
    'cn': 'notes/moc cn.md',
    'algorithms': 'notes/moc algo.md',
    'programming and ds': 'notes/moc dsa.md',
    'data structures': 'notes/moc dsa.md',
    'c programming': 'notes/moc clang.md',
    'c-programming': 'notes/moc clang.md',
    'engineering mathematics': 'notes/moc maths.md',
    'general aptitude': 'notes/moc aptitude.md',
    'aptitude': 'notes/moc aptitude.md',
    'calculus': 'notes/moc calculus.md',
    'linear algebra': 'notes/moc la.md',
    'probability': 'notes/moc probablity.md',
    'python': 'notes/moc python.md',
}

created_count = 0

for sheet_name, rel_folder in folder_map.items():
    if sheet_name not in wb.sheetnames:
        continue
    
    out_dir = os.path.join(VAULT_PATH, rel_folder)
    os.makedirs(out_dir, exist_ok=True)
    
    sheet = wb[sheet_name]
    
    for r in range(2, sheet.max_row + 1):
        if sheet_name == 'CSE quiz':
            done = bool(sheet.cell(r, 1).value)
            subj = sheet.cell(r, 3).value
            tname = sheet.cell(r, 4).value
            qs = sheet.cell(r, 5).value
            marks = sheet.cell(r, 6).value
            duration = sheet.cell(r, 7).value
            taken = sheet.cell(r, 8).value
            url = sheet.cell(r, 9).value
        else:
            done = bool(sheet.cell(r, 1).value)
            subj = sheet.cell(r, 2).value
            tname = sheet.cell(r, 3).value
            qs = sheet.cell(r, 4).value
            marks = sheet.cell(r, 5).value
            duration = sheet.cell(r, 6).value
            taken = sheet.cell(r, 7).value
            url = ""

        if not tname and not subj:
            continue
            
        subj_str = str(subj).strip() if subj else "General"
        tname_str = str(tname).strip() if tname else f"Test {r-1}"
        
        safe_tname = clean_filename(tname_str)
        filename = f"{r-1:03d} - {safe_tname}.md"
        filepath = os.path.join(out_dir, filename)
        
        # Link to MOC
        moc_file = subject_moc_map.get(subj_str.lower(), "")
        if moc_file:
            moc_link = f"[[{moc_file[:-3]}|{subj_str}]]"
        else:
            moc_link = subj_str

        status_str = "completed" if done else "pending"
        date_comp = "2026-09-25" if done else ""

        content = f"""---
test_id: "{sheet_name[:3].upper()}-{r-1:04d}"
category: "{sheet_name}"
subject: "{subj_str}"
test_name: "{tname_str}"
questions: "{qs or ''}"
marks: "{marks or ''}"
duration: "{duration or ''}"
taken_count: {taken or 0}
status: "{status_str}"
date_completed: "{date_comp}"
score_obtained: 
weak_areas: []
url: "{url or ''}"
---

# {tname_str}

## 📌 Test Details
- **Category**: {sheet_name}
- **Subject**: {moc_link}
- **Questions**: {qs or 'N/A'} | **Marks**: {marks or 'N/A'} | **Duration**: {duration or 'N/A'}
- **Status**: `{status_str}`

---

## 📝 Analysis & Mistakes
- [ ] Review key concepts and wrong attempts from this test.

## 💡 Variation Questions & Practice Notes
- Add custom variation questions or detailed revision notes here.
"""
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)
        created_count += 1

print(f"Successfully generated {created_count} test notes in {VAULT_PATH}/Tests/")
